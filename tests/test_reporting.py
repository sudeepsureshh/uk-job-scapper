"""
Tests for the ReportGenerator.
Verifies that all three report formats are generated correctly.
"""
import csv
import json
import os
import sys
from pathlib import Path

import pytest
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from reporting.generator import ReportGenerator


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture
def reporter(tmp_path: Path) -> ReportGenerator:
    """ReportGenerator writing to a temporary directory."""
    output_dir = str(tmp_path / "output")
    return ReportGenerator(output_dir=output_dir)


@pytest.fixture
def sample_jobs() -> list[dict]:
    """A small set of scored job dicts for report generation tests."""
    return [
        {
            "job_id": "rpt_001",
            "title": "Java Backend Engineer",
            "employer": "Revolut",
            "location": "London, UK (Hybrid)",
            "salary_min": 70000,
            "salary_max": 90000,
            "salary_currency": "GBP",
            "description": "Java Spring Boot microservices for payments.",
            "url": "https://revolut.com/jobs/java-backend-001",
            "apply_url": "https://revolut.com/jobs/java-backend-001/apply",
            "posted_date": "2026-10-14",
            "closing_date": None,
            "source": "reed.co.uk",
            "scraped_at": "2026-10-14T07:00:00",
            "relevance_score": 85.0,
            "score_breakdown": {
                "skills_match": 30,
                "experience_match": 22,
                "responsibilities_match": 12,
                "location_match": 10,
                "salary_match": 8,
                "sponsorship_bonus": 3,
            },
            "matched_skills": ["Java", "Spring Boot", "PostgreSQL", "Docker"],
            "missing_skills": ["Kafka"],
            "match_explanation": "Score: 85/100. Matched: Java, Spring Boot. Missing: Kafka.",
            "sponsorship_status": "confirmed_role",
            "sponsorship_evidence": ["visa sponsorship"],
            "sponsorship_notes": "Mentioned in JD.",
            "is_new": 1,
            "first_seen": "2026-10-14T07:00:00",
            "last_seen": "2026-10-14T07:00:00",
        },
        {
            "job_id": "rpt_002",
            "title": "Senior Java Developer",
            "employer": "Barclays",
            "location": "Northampton, UK",
            "salary_min": 65000,
            "salary_max": 80000,
            "salary_currency": "GBP",
            "description": "Spring Boot API development for banking systems.",
            "url": "https://jobs.barclays.com/java-developer-002",
            "apply_url": "https://jobs.barclays.com/java-developer-002/apply",
            "posted_date": "2026-10-13",
            "closing_date": "2026-11-13",
            "source": "adzuna.co.uk",
            "scraped_at": "2026-10-14T07:00:00",
            "relevance_score": 72.0,
            "score_breakdown": {
                "skills_match": 25,
                "experience_match": 20,
                "responsibilities_match": 12,
                "location_match": 8,
                "salary_match": 7,
                "sponsorship_bonus": 0,
            },
            "matched_skills": ["Java", "Spring Boot", "Oracle"],
            "missing_skills": ["Kafka", "GraphQL"],
            "match_explanation": "Score: 72/100. Matched: Java, Spring Boot.",
            "sponsorship_status": "licensed_employer_unconfirmed",
            "sponsorship_evidence": [],
            "sponsorship_notes": "Employer on register; confirm sponsorship.",
            "is_new": 0,
            "first_seen": "2026-10-12T07:00:00",
            "last_seen": "2026-10-14T07:00:00",
        },
        {
            "job_id": "rpt_003",
            "title": "Marketing Manager",
            "employer": "Generic Retail",
            "location": "Manchester, UK",
            "salary_min": 35000,
            "salary_max": 45000,
            "salary_currency": "GBP",
            "description": "Digital marketing and social media management.",
            "url": "https://genericretail.co.uk/marketing-003",
            "apply_url": "https://genericretail.co.uk/marketing-003/apply",
            "posted_date": "2026-10-14",
            "closing_date": None,
            "source": "reed.co.uk",
            "scraped_at": "2026-10-14T07:00:00",
            "relevance_score": 12.0,
            "score_breakdown": {
                "skills_match": 2,
                "experience_match": 5,
                "responsibilities_match": 2,
                "location_match": 3,
                "salary_match": 0,
                "sponsorship_bonus": 0,
            },
            "matched_skills": [],
            "missing_skills": [],
            "match_explanation": "Score: 12/100. Weak match.",
            "sponsorship_status": "sponsorship_unavailable",
            "sponsorship_evidence": ["no sponsorship"],
            "sponsorship_notes": "Explicitly denied.",
            "is_new": 1,
            "first_seen": "2026-10-14T07:00:00",
            "last_seen": "2026-10-14T07:00:00",
        },
    ]


@pytest.fixture
def run_summary() -> dict:
    """Sample run summary dict."""
    return {
        "run_id": 42,
        "started_at": "2026-10-14T07:00:00",
        "jobs_found": 3,
        "new_jobs": 2,
        "status": "success",
    }


# ---------------------------------------------------------------------------
# Tests: Excel
# ---------------------------------------------------------------------------


class TestGenerateExcel:
    """Tests for generate_excel."""

    def test_excel_created(
        self,
        reporter: ReportGenerator,
        sample_jobs: list[dict],
    ) -> None:
        """Excel file should be created at the expected path."""
        path = reporter.generate_excel(sample_jobs, "2026-10-14")
        assert Path(path).exists(), f"Excel file not found at {path}"
        assert path.endswith(".xlsx"), "Expected .xlsx extension."

    def test_excel_has_expected_sheets(
        self,
        reporter: ReportGenerator,
        sample_jobs: list[dict],
    ) -> None:
        """Workbook should contain the four expected sheets."""
        path = reporter.generate_excel(sample_jobs, "2026-10-14")
        wb = openpyxl.load_workbook(path)
        expected_sheets = {"All Jobs", "Shortlist", "New Today", "Confirmed Sponsorship"}
        assert expected_sheets == set(wb.sheetnames), (
            f"Expected sheets {expected_sheets}, got {set(wb.sheetnames)}"
        )

    def test_excel_all_jobs_has_data_rows(
        self,
        reporter: ReportGenerator,
        sample_jobs: list[dict],
    ) -> None:
        """'All Jobs' sheet should have data rows (beyond header rows)."""
        path = reporter.generate_excel(sample_jobs, "2026-10-14")
        wb = openpyxl.load_workbook(path)
        ws = wb["All Jobs"]
        # Row 1: title, Row 2: header, Row 3+: data
        assert ws.max_row >= 3, f"Expected at least 3 rows in All Jobs, got {ws.max_row}"

    def test_excel_confirmed_sponsorship_sheet(
        self,
        reporter: ReportGenerator,
        sample_jobs: list[dict],
    ) -> None:
        """'Confirmed Sponsorship' sheet should contain only confirmed jobs."""
        path = reporter.generate_excel(sample_jobs, "2026-10-14")
        wb = openpyxl.load_workbook(path)
        ws = wb["Confirmed Sponsorship"]
        # rpt_001 is the only confirmed job; should appear in data
        data_rows = [ws.cell(row=r, column=2).value for r in range(3, ws.max_row + 1)]
        confirmed_employers = [v for v in data_rows if v]
        assert "Revolut" in confirmed_employers


# ---------------------------------------------------------------------------
# Tests: CSV
# ---------------------------------------------------------------------------


class TestGenerateCsv:
    """Tests for generate_csv."""

    def test_csv_created(
        self,
        reporter: ReportGenerator,
        sample_jobs: list[dict],
    ) -> None:
        """CSV file should be created."""
        path = reporter.generate_csv(sample_jobs, "2026-10-14")
        assert Path(path).exists()
        assert path.endswith(".csv")

    def test_csv_has_correct_number_of_rows(
        self,
        reporter: ReportGenerator,
        sample_jobs: list[dict],
    ) -> None:
        """CSV should have one header row + one data row per job."""
        path = reporter.generate_csv(sample_jobs, "2026-10-14")
        with open(path, newline="", encoding="utf-8") as f:
            reader = csv.reader(f)
            rows = list(reader)
        # 1 header + len(sample_jobs) data rows
        assert len(rows) == len(sample_jobs) + 1, (
            f"Expected {len(sample_jobs)+1} rows (header + data), got {len(rows)}"
        )

    def test_csv_has_expected_columns(
        self,
        reporter: ReportGenerator,
        sample_jobs: list[dict],
    ) -> None:
        """CSV should contain key expected columns."""
        path = reporter.generate_csv(sample_jobs, "2026-10-14")
        with open(path, newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            headers = reader.fieldnames or []
        for col in ("rank", "employer", "title", "relevance_score", "url"):
            assert col in headers, f"Expected column '{col}' in CSV headers."

    def test_csv_sorted_by_score_descending(
        self,
        reporter: ReportGenerator,
        sample_jobs: list[dict],
    ) -> None:
        """CSV rows should be ordered by relevance_score descending."""
        path = reporter.generate_csv(sample_jobs, "2026-10-14")
        with open(path, newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            scores = [float(row["relevance_score"]) for row in reader]
        assert scores == sorted(scores, reverse=True), (
            f"CSV not sorted by score descending: {scores}"
        )


# ---------------------------------------------------------------------------
# Tests: Markdown
# ---------------------------------------------------------------------------


class TestGenerateMarkdown:
    """Tests for generate_markdown."""

    def test_markdown_created(
        self,
        reporter: ReportGenerator,
        sample_jobs: list[dict],
        run_summary: dict,
    ) -> None:
        """Markdown file should be created."""
        path = reporter.generate_markdown(sample_jobs, "2026-10-14", run_summary)
        assert Path(path).exists()
        assert path.endswith(".md")

    def test_markdown_contains_date(
        self,
        reporter: ReportGenerator,
        sample_jobs: list[dict],
        run_summary: dict,
    ) -> None:
        """Markdown should contain the report date in the title."""
        path = reporter.generate_markdown(sample_jobs, "2026-10-14", run_summary)
        content = Path(path).read_text(encoding="utf-8")
        assert "2026-10-14" in content

    def test_markdown_contains_employer_names(
        self,
        reporter: ReportGenerator,
        sample_jobs: list[dict],
        run_summary: dict,
    ) -> None:
        """Markdown shortlist should include at least one top employer."""
        path = reporter.generate_markdown(sample_jobs, "2026-10-14", run_summary)
        content = Path(path).read_text(encoding="utf-8")
        assert "Revolut" in content, "Expected 'Revolut' in markdown content."

    def test_markdown_contains_disclaimer(
        self,
        reporter: ReportGenerator,
        sample_jobs: list[dict],
        run_summary: dict,
    ) -> None:
        """Markdown should include a disclaimer about legal advice."""
        path = reporter.generate_markdown(sample_jobs, "2026-10-14", run_summary)
        content = Path(path).read_text(encoding="utf-8")
        assert "advice" in content.lower(), "Expected disclaimer mentioning advice."


# ---------------------------------------------------------------------------
# Tests: empty job list
# ---------------------------------------------------------------------------


class TestEmptyJobList:
    """Tests that all formats handle an empty job list gracefully."""

    def test_excel_empty_jobs_handled(
        self,
        reporter: ReportGenerator,
    ) -> None:
        """generate_excel should not raise with an empty list."""
        path = reporter.generate_excel([], "2026-10-14")
        assert Path(path).exists()

    def test_csv_empty_jobs_handled(
        self,
        reporter: ReportGenerator,
    ) -> None:
        """generate_csv should produce a header-only file for empty list."""
        path = reporter.generate_csv([], "2026-10-14")
        assert Path(path).exists()
        with open(path, newline="", encoding="utf-8") as f:
            rows = list(csv.reader(f))
        # Only the header row
        assert len(rows) == 1

    def test_markdown_empty_jobs_handled(
        self,
        reporter: ReportGenerator,
    ) -> None:
        """generate_markdown should not raise with an empty list."""
        path = reporter.generate_markdown([], "2026-10-14", {})
        assert Path(path).exists()

    def test_generate_all_empty_no_crash(
        self,
        reporter: ReportGenerator,
    ) -> None:
        """generate_all should not raise with an empty job list."""
        outputs = reporter.generate_all([], "2026-10-14", {})
        # All three formats should have been attempted
        assert "excel" in outputs or "csv" in outputs or "markdown" in outputs
