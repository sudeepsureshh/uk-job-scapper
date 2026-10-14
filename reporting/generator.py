"""
Report generator: produces Excel, CSV, and Markdown reports from job data.
"""
import csv
import json
import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

import openpyxl
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Styling constants
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill(
    start_color="1F4E79", end_color="1F4E79", fill_type="solid"
)
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)

ALT_ROW_FILL = PatternFill(
    start_color="DCE6F1", end_color="DCE6F1", fill_type="solid"
)

CONFIRMED_FILL = PatternFill(
    start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

# Columns for the main jobs sheet
COLUMNS = [
    ("Rank", 6),
    ("Employer", 28),
    ("Title", 40),
    ("Location", 20),
    ("Salary", 22),
    ("Score", 8),
    ("Sponsorship", 22),
    ("Sponsorship Evidence", 35),
    ("Matched Skills", 40),
    ("Missing Skills", 30),
    ("Match Explanation", 55),
    ("Posted Date", 14),
    ("Closing Date", 14),
    ("URL", 50),
    ("Apply URL", 50),
]


class ReportGenerator:
    """
    Generates Excel, CSV, and Markdown reports from a list of scored jobs.

    All output is written to the configured output directory.
    Handles empty job lists gracefully.
    """

    def __init__(self, output_dir: str = "output") -> None:
        """
        Initialise the generator.

        Args:
            output_dir: Directory where reports will be written.
        """
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate_all(
        self,
        jobs: list[dict],
        date_str: str,
        run_summary: dict,
    ) -> dict[str, str]:
        """
        Generate all three report formats.

        Args:
            jobs: List of scored job dicts.
            date_str: Date string for filenames (e.g. "2026-10-14").
            run_summary: Dict with run statistics.

        Returns:
            Dict mapping format names to output file paths.
        """
        output = {}

        try:
            output["excel"] = self.generate_excel(jobs, date_str)
            logger.info("Excel report: %s", output["excel"])
        except Exception as exc:
            logger.error("Failed to generate Excel report: %s", exc)

        try:
            output["csv"] = self.generate_csv(jobs, date_str)
            logger.info("CSV report: %s", output["csv"])
        except Exception as exc:
            logger.error("Failed to generate CSV report: %s", exc)

        try:
            output["markdown"] = self.generate_markdown(jobs, date_str, run_summary)
            logger.info("Markdown report: %s", output["markdown"])
        except Exception as exc:
            logger.error("Failed to generate Markdown report: %s", exc)

        return output

    # ------------------------------------------------------------------
    # Excel report
    # ------------------------------------------------------------------

    def generate_excel(self, jobs: list[dict], date_str: str) -> str:
        """
        Generate a formatted Excel workbook.

        Sheets:
          1. All Jobs
          2. Shortlist (top 10 by relevance score)
          3. New Today (jobs with is_new=True)
          4. Confirmed Sponsorship

        Args:
            jobs: List of scored job dicts.
            date_str: Date string for filename.

        Returns:
            Path to the generated .xlsx file.
        """
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default empty sheet

        sorted_jobs = sorted(jobs, key=lambda j: j.get("relevance_score") or 0, reverse=True)
        shortlist = sorted_jobs[:10]
        new_today = [j for j in sorted_jobs if j.get("is_new")]
        confirmed = [j for j in sorted_jobs if j.get("sponsorship_status") == "confirmed_role"]

        # Sheet 1: All Jobs
        ws_all = wb.create_sheet("All Jobs")
        self._write_jobs_sheet(ws_all, sorted_jobs, "All Jobs")

        # Sheet 2: Shortlist
        ws_short = wb.create_sheet("Shortlist")
        self._write_jobs_sheet(ws_short, shortlist, "Top 10 by Relevance Score")

        # Sheet 3: New Today
        ws_new = wb.create_sheet("New Today")
        self._write_jobs_sheet(ws_new, new_today, "New Jobs Discovered Today")

        # Sheet 4: Confirmed Sponsorship
        ws_sponsor = wb.create_sheet("Confirmed Sponsorship")
        self._write_jobs_sheet(ws_sponsor, confirmed, "Confirmed Sponsorship Available")

        filepath = self.output_dir / f"jobs_{date_str}.xlsx"
        wb.save(str(filepath))
        return str(filepath)

    def _write_jobs_sheet(
        self,
        ws: Worksheet,
        jobs: list[dict],
        title: str,
    ) -> None:
        """Write a list of jobs to a worksheet with formatting."""
        # Title row
        ws.append([title])
        title_cell = ws.cell(row=1, column=1)
        title_cell.font = Font(bold=True, size=14, color="1F4E79")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))

        # Header row
        headers = [col[0] for col in COLUMNS]
        ws.append(headers)
        header_row = ws.max_row

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", wrap_text=False)

        # Data rows
        for rank, job in enumerate(jobs, start=1):
            row_data = self._job_to_row(job, rank)
            ws.append(row_data)
            data_row = ws.max_row

            # Alternate row shading
            fill = ALT_ROW_FILL if rank % 2 == 0 else None

            # Sponsorship highlight
            sponsor_status = job.get("sponsorship_status", "")
            if sponsor_status == "confirmed_role":
                fill = CONFIRMED_FILL

            if fill:
                for col_idx in range(1, len(COLUMNS) + 1):
                    ws.cell(row=data_row, column=col_idx).fill = fill

            # Hyperlink URL columns (14 and 15)
            url = job.get("url") or ""
            apply_url = job.get("apply_url") or ""
            if url:
                url_cell = ws.cell(row=data_row, column=14)
                url_cell.hyperlink = url
                url_cell.font = Font(color="0563C1", underline="single")
            if apply_url:
                apply_cell = ws.cell(row=data_row, column=15)
                apply_cell.hyperlink = apply_url
                apply_cell.font = Font(color="0563C1", underline="single")

        # Freeze top two rows (title + header)
        ws.freeze_panes = "A3"

        # Set column widths
        for col_idx, (_, width) in enumerate(COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # Auto-filter on header row
        if ws.max_row >= 2:
            ws.auto_filter.ref = (
                f"A{header_row}:{get_column_letter(len(COLUMNS))}{ws.max_row}"
            )

        # Empty state message
        if not jobs:
            ws.cell(row=3, column=1).value = "No jobs found for this filter."
            ws.cell(row=3, column=1).font = Font(italic=True, color="808080")

    def _job_to_row(self, job: dict, rank: int) -> list[Any]:
        """Convert a job dict to a spreadsheet row."""
        # Format salary
        sal_min = job.get("salary_min")
        sal_max = job.get("salary_max")
        currency = job.get("salary_currency", "GBP")
        symbol = "£" if currency == "GBP" else currency + " "
        if sal_min and sal_max and sal_min != sal_max:
            salary_str = f"{symbol}{int(sal_min):,} - {symbol}{int(sal_max):,}"
        elif sal_min or sal_max:
            val = sal_max or sal_min
            salary_str = f"{symbol}{int(val):,}"  # type: ignore
        else:
            salary_str = "Not stated"

        # Format skills
        matched = job.get("matched_skills") or []
        if isinstance(matched, str):
            try:
                matched = json.loads(matched)
            except Exception:
                matched = [matched]
        matched_str = ", ".join(matched[:10])

        missing = job.get("missing_skills") or []
        if isinstance(missing, str):
            try:
                missing = json.loads(missing)
            except Exception:
                missing = [missing]
        missing_str = ", ".join(missing[:8])

        # Format sponsorship evidence
        evidence = job.get("sponsorship_evidence") or []
        if isinstance(evidence, str):
            try:
                evidence = json.loads(evidence)
            except Exception:
                evidence = [evidence]
        evidence_str = ", ".join(evidence[:3])

        # Sponsorship display text
        status = job.get("sponsorship_status", "")
        status_display = {
            "confirmed_role": "Confirmed",
            "licensed_employer_unconfirmed": "Likely (Licensed Employer)",
            "sponsorship_unavailable": "Not Available",
            "unable_to_verify": "Unverified",
        }.get(status, status or "Unknown")

        return [
            rank,
            job.get("employer", ""),
            job.get("title", ""),
            job.get("location", ""),
            salary_str,
            round(job.get("relevance_score") or 0, 1),
            status_display,
            evidence_str,
            matched_str,
            missing_str,
            job.get("match_explanation", ""),
            job.get("posted_date", ""),
            job.get("closing_date", ""),
            job.get("url", ""),
            job.get("apply_url", ""),
        ]

    # ------------------------------------------------------------------
    # CSV report
    # ------------------------------------------------------------------

    def generate_csv(self, jobs: list[dict], date_str: str) -> str:
        """
        Generate a CSV report.

        Args:
            jobs: List of scored job dicts.
            date_str: Date string for filename.

        Returns:
            Path to the generated .csv file.
        """
        filepath = self.output_dir / f"jobs_{date_str}.csv"
        fieldnames = [
            "rank", "employer", "title", "location", "salary",
            "relevance_score", "sponsorship_status", "sponsorship_evidence",
            "matched_skills", "missing_skills", "match_explanation",
            "posted_date", "closing_date", "url", "apply_url",
            "source", "is_new", "first_seen",
        ]

        sorted_jobs = sorted(jobs, key=lambda j: j.get("relevance_score") or 0, reverse=True)

        with open(filepath, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()

            for rank, job in enumerate(sorted_jobs, start=1):
                sal_min = job.get("salary_min")
                sal_max = job.get("salary_max")
                if sal_min and sal_max and sal_min != sal_max:
                    salary = f"£{int(sal_min):,} - £{int(sal_max):,}"
                elif sal_min or sal_max:
                    val = sal_max or sal_min
                    salary = f"£{int(val):,}"  # type: ignore
                else:
                    salary = ""

                matched = job.get("matched_skills") or []
                if isinstance(matched, str):
                    try:
                        matched = json.loads(matched)
                    except Exception:
                        matched = [matched]

                missing = job.get("missing_skills") or []
                if isinstance(missing, str):
                    try:
                        missing = json.loads(missing)
                    except Exception:
                        missing = [missing]

                evidence = job.get("sponsorship_evidence") or []
                if isinstance(evidence, str):
                    try:
                        evidence = json.loads(evidence)
                    except Exception:
                        evidence = [evidence]

                writer.writerow({
                    "rank": rank,
                    "employer": job.get("employer", ""),
                    "title": job.get("title", ""),
                    "location": job.get("location", ""),
                    "salary": salary,
                    "relevance_score": round(job.get("relevance_score") or 0, 1),
                    "sponsorship_status": job.get("sponsorship_status", ""),
                    "sponsorship_evidence": "; ".join(evidence[:3]),
                    "matched_skills": ", ".join(matched[:10]),
                    "missing_skills": ", ".join(missing[:8]),
                    "match_explanation": job.get("match_explanation", ""),
                    "posted_date": job.get("posted_date", ""),
                    "closing_date": job.get("closing_date", ""),
                    "url": job.get("url", ""),
                    "apply_url": job.get("apply_url", ""),
                    "source": job.get("source", ""),
                    "is_new": job.get("is_new", 0),
                    "first_seen": job.get("first_seen", ""),
                })

        return str(filepath)

    # ------------------------------------------------------------------
    # Markdown report
    # ------------------------------------------------------------------

    def generate_markdown(
        self,
        jobs: list[dict],
        date_str: str,
        run_summary: dict,
    ) -> str:
        """
        Generate a Markdown summary report.

        Args:
            jobs: List of scored job dicts.
            date_str: Date string for filename.
            run_summary: Dict with run statistics.

        Returns:
            Path to the generated .md file.
        """
        filepath = self.output_dir / f"jobs_{date_str}.md"
        sorted_jobs = sorted(jobs, key=lambda j: j.get("relevance_score") or 0, reverse=True)
        shortlist = sorted_jobs[:10]
        new_today = [j for j in sorted_jobs if j.get("is_new")]
        confirmed = [j for j in sorted_jobs if j.get("sponsorship_status") == "confirmed_role"]

        lines: list[str] = []

        lines.append(f"# UK Job Search Report — {date_str}")
        lines.append("")
        lines.append("> DISCLAIMER: This report is generated automatically. "
                     "Always verify job details and sponsorship directly with employers. "
                     "This does not constitute legal or immigration advice.")
        lines.append("")

        # Summary stats
        lines.append("## Summary")
        lines.append("")
        lines.append(f"| Metric | Value |")
        lines.append(f"|--------|-------|")
        lines.append(f"| Total jobs found | {len(jobs)} |")
        lines.append(f"| New today | {len(new_today)} |")
        lines.append(f"| With confirmed sponsorship | {len(confirmed)} |")
        lines.append(f"| Report generated | {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')} |")

        if run_summary:
            lines.append(f"| Scrape run started | {run_summary.get('started_at', 'N/A')} |")
            lines.append(f"| Scrape run status | {run_summary.get('status', 'N/A')} |")

        lines.append("")

        # Shortlist
        lines.append("## Top 10 by Relevance Score")
        lines.append("")
        if shortlist:
            lines.append("| # | Employer | Title | Location | Score | Sponsorship | URL |")
            lines.append("|---|----------|-------|----------|-------|-------------|-----|")
            for i, job in enumerate(shortlist, start=1):
                employer = _md_escape(job.get("employer", ""))
                title = _md_escape(job.get("title", ""))
                location = _md_escape(job.get("location", ""))
                score = round(job.get("relevance_score") or 0, 0)
                status = _sponsorship_badge(job.get("sponsorship_status", ""))
                url = job.get("url") or ""
                url_link = f"[Apply]({url})" if url else "N/A"
                lines.append(f"| {i} | {employer} | {title} | {location} | {score} | {status} | {url_link} |")
        else:
            lines.append("No jobs found.")

        lines.append("")

        # New Today section
        lines.append("## Newly Discovered Jobs")
        lines.append("")
        if new_today:
            lines.append(f"Found **{len(new_today)}** new jobs today.\n")
            for job in new_today[:20]:
                title = job.get("title", "Unknown")
                employer = job.get("employer", "Unknown")
                score = round(job.get("relevance_score") or 0, 0)
                url = job.get("url") or ""
                status = _sponsorship_badge(job.get("sponsorship_status", ""))
                explanation = job.get("match_explanation", "")
                lines.append(f"### {employer} — {title}")
                lines.append(f"**Score:** {score}/100 | **Sponsorship:** {status}")
                if url:
                    lines.append(f"**URL:** {url}")
                if explanation:
                    lines.append(f"**Notes:** {explanation}")
                lines.append("")
        else:
            lines.append("No new jobs found in this run.")
            lines.append("")

        # Confirmed sponsorship section
        lines.append("## Confirmed Sponsorship Available")
        lines.append("")
        if confirmed:
            for job in confirmed[:15]:
                title = job.get("title", "Unknown")
                employer = job.get("employer", "Unknown")
                score = round(job.get("relevance_score") or 0, 0)
                url = job.get("url") or ""
                evidence = job.get("sponsorship_evidence") or []
                if isinstance(evidence, str):
                    try:
                        evidence = json.loads(evidence)
                    except Exception:
                        evidence = [evidence]
                evidence_str = ", ".join(evidence[:2]) if evidence else ""
                lines.append(f"- **{employer}** — {title} (Score: {score}) {f'[Apply]({url})' if url else ''}")
                if evidence_str:
                    lines.append(f"  Evidence: _{evidence_str}_")
        else:
            lines.append("No jobs with confirmed sponsorship found in this run.")

        lines.append("")
        lines.append("---")
        lines.append("*Generated by UK Job Scraper — https://github.com/your-username/uk-job-scraper*")
        lines.append("*This report does not constitute legal or immigration advice.*")

        content = "\n".join(lines)
        filepath.write_text(content, encoding="utf-8")
        return str(filepath)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _md_escape(text: str) -> str:
    """Escape pipe characters for Markdown table cells."""
    return text.replace("|", "\\|").replace("\n", " ").strip()


def _sponsorship_badge(status: str) -> str:
    """Return a human-readable sponsorship status badge."""
    return {
        "confirmed_role": "Confirmed",
        "licensed_employer_unconfirmed": "Likely",
        "sponsorship_unavailable": "Not Available",
        "unable_to_verify": "Unverified",
    }.get(status, "Unknown")
